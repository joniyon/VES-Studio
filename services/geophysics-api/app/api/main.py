"""Thin HTTP adapter over the engine. No science lives here."""
import io
import math
from dataclasses import asdict

import pandas as pd
from fastapi import FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field

from app import ENGINE_VERSION
from app.arrays import REGISTRY, get_array
from app.interpretation import CATALOGUE, CONFIDENCE_LEVELS, get_lithology, suggest
from app.inversion import InversionConfig, InversionError, invert_station
from app.processing import process_dataset
from app.reports import (build_report, column_figure, curve_figure, draft_summary, layers_csv, legend_figure,
                         processed_csv)
from app.units import CURRENT, DISTANCE, RESISTANCE, RESISTIVITY, VOLTAGE

app = FastAPI(title="VES Studio geophysics API", version=ENGINE_VERSION)
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
                   allow_methods=["*"], allow_headers=["*"])

FIELD_LABELS = {
    "ab_half": "AB/2", "mn_half": "MN/2", "a": "a (spacing)", "n": "n (separation factor)",
    "ab": "AB", "mn": "MN", "x": "x (MN centre offset)",
}
ARRAY_TITLES = {
    "schlumberger": "Schlumberger", "wenner": "Wenner", "dipole_dipole": "Dipole-dipole",
    "pole_dipole": "Pole-dipole", "pole_pole": "Pole-pole", "gradient": "Gradient / rectangle",
}
MEDIA = {"png": "image/png", "svg": "image/svg+xml"}


def clean(v):
    """Make engine output JSON-safe: NaN/inf -> None, numpy -> python."""
    if isinstance(v, dict):
        return {k: clean(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [clean(x) for x in v]
    if hasattr(v, "tolist"):
        return clean(v.tolist())
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


class Dataset(BaseModel):
    array: str
    units: dict[str, str] = Field(default_factory=dict)
    rows: list[dict]
    excluded_rows: list[int] = Field(default_factory=list)
    assume_point_mn: bool = False


class InvertRequest(Dataset):
    config: dict = Field(default_factory=dict)


class FigureRequest(InvertRequest):
    interpretation: list[dict] = Field(default_factory=list)
    with_model: bool = True


class ReportRequest(InvertRequest):
    project: dict = Field(default_factory=dict)
    station: dict = Field(default_factory=dict)
    interpretation: list[dict] = Field(default_factory=list)
    geological_context: str = ""
    conclusion: str = ""
    upload: dict = Field(default_factory=dict)


class LayerIn(BaseModel):
    resistivity: float
    depth_top: float
    depth_bottom: float | None = None


class SuggestRequest(BaseModel):
    layers: list[LayerIn]


@app.exception_handler(ValueError)
async def value_error(_, exc: ValueError):
    return JSONResponse({"error": str(exc)}, status_code=422)


@app.get("/health")
def health():
    return {"status": "ok", "engine_version": ENGINE_VERSION}


@app.get("/arrays")
def arrays():
    return {
        "arrays": [
            {"id": k, "title": ARRAY_TITLES[k],
             "fields": [{"name": f, "label": FIELD_LABELS[f], "kind": "distance"} for f in a.required_fields]
                       + [{"name": f, "label": FIELD_LABELS[f], "kind": "integer"} for f in a.integer_fields],
             "references": list(a.references)}
            for k, a in REGISTRY.items()
        ],
        "units": {"distance": list(DISTANCE), "resistance": list(RESISTANCE),
                  "voltage": list(VOLTAGE), "current": list(CURRENT),
                  "resistivity": list(RESISTIVITY)},
    }


def _process(req: Dataset):
    get_array(req.array)  # raises ValueError -> 422
    return process_dataset(pd.DataFrame(req.rows), req.array, req.units, set(req.excluded_rows), req.assume_point_mn)


def _config(req: InvertRequest) -> InversionConfig:
    return InversionConfig(**{k: (tuple(v) if isinstance(v, list) else v) for k, v in req.config.items()})


def _run(req: InvertRequest):
    processed = _process(req)
    if processed.table.empty:
        raise InversionError("; ".join(i.message for i in processed.issues))
    return processed, invert_station(processed, _config(req))


@app.post("/process")
def process(req: Dataset):
    res = _process(req)
    return clean({
        "rows": res.table.to_dict(orient="records"),
        "issues": [{**asdict(i), "severity": i.severity.value} for i in res.issues],
        "has_errors": res.has_errors,
        "lineage": res.lineage,
    })


@app.post("/invert")
def invert(req: InvertRequest):
    try:
        _, r = _run(req)
    except InversionError as e:
        return clean({"error": str(e)})
    d = asdict(r)
    d["config"] = asdict(r.config)
    return clean(d)


# ---- interpretation -------------------------------------------------------------------------
@app.get("/lithologies")
def lithologies():
    return {
        "lithologies": [{"id": l.id, "name": l.name, "group": l.group, "color": l.color, "hatch": l.hatch,
                         "rho_min": l.rho_min, "rho_max": l.rho_max, "verified": l.verified} for l in CATALOGUE],
        "confidence_levels": list(CONFIDENCE_LEVELS),
        "note": "Resistivity ranges are indicative and overlapping; they do not identify lithology.",
    }


@app.post("/suggest-lithology")
def suggest_lithology(req: SuggestRequest):
    return clean({"layers": [
        [{"id": s.lithology.id, "name": s.lithology.name, "in_range": s.in_range, "depth_ok": s.depth_ok,
          "score": s.score, "basis": s.basis} for s in suggest(l.resistivity, l.depth_top, l.depth_bottom)]
        for l in req.layers]})


# ---- file parsing -----------------------------------------------------------------------------
@app.post("/parse-xlsx")
async def parse_xlsx(file: UploadFile = File(...), sheet: str | None = Form(None)):
    from openpyxl import load_workbook
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("File is larger than 10 MB.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Could not read this file as an .xlsx workbook.") from None
    name = sheet if sheet in wb.sheetnames else wb.sheetnames[0]
    rows = [["" if c is None else str(c) for c in r] for r in wb[name].iter_rows(values_only=True)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ValueError(f"Sheet '{name}' is empty.")
    return {"sheets": wb.sheetnames, "sheet": name, "headers": [h.strip() for h in rows[0]], "rows": rows[1:]}


# ---- figures, exports, report -------------------------------------------------------------------
@app.post("/figures/curve")
def fig_curve(req: FigureRequest, fmt: str = "png"):
    if req.with_model:
        processed, r = _run(req)
    else:
        processed, r = _process(req), None
    return Response(curve_figure(processed.table, r, fmt), media_type=MEDIA.get(fmt, "image/png"))


@app.post("/figures/column")
def fig_column(req: FigureRequest, fmt: str = "png"):
    _, r = _run(req)
    return Response(column_figure(r, req.interpretation, fmt), media_type=MEDIA.get(fmt, "image/png"))


@app.post("/export/processed.csv")
def export_processed(req: Dataset):
    return Response(processed_csv(_process(req)), media_type="text/csv")


@app.post("/export/layers.csv")
def export_layers(req: FigureRequest):
    _, r = _run(req)
    return Response(layers_csv(r, req.interpretation), media_type="text/csv")


@app.post("/report.pdf")
def report(req: ReportRequest):
    processed, r = _run(req)
    pdf = build_report(project=req.project, station=req.station, processed=processed, result=r,
                       interpretation=req.interpretation, geological_context=req.geological_context,
                       conclusion=req.conclusion, upload=req.upload)
    return Response(pdf, media_type="application/pdf")


@app.post("/draft-summary")
def draft(req: FigureRequest):
    processed, r = _run(req)
    return {"summary": draft_summary(processed, r, req.interpretation)}


@app.get("/figures/legend")
def fig_legend(ids: str = "", fmt: str = "png"):
    """Lithology legend; `ids` is a comma-separated list (default: the whole catalogue)."""
    chosen = [i for i in ids.split(",") if i] or [l.id for l in CATALOGUE]
    for i in chosen:
        get_lithology(i)
    return Response(legend_figure(chosen, fmt), media_type=MEDIA.get(fmt, "image/png"))
